import os
import shutil
import datetime
import math
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter
from config import Config
from shared.thai_date import to_thai_date_excel_short
from shared.date_utils import to_thai_date_full

thai_months = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"
]

def _set_cell_value(ws, coord, val):
    """Sets value on cell only if it's not a MergedCell."""
    cell = ws[coord]
    if cell.__class__.__name__ != 'MergedCell':
        cell.value = val

def _set_cell_value_rc(ws, r, c, val):
    """Sets value on cell by row/col only if it's not a MergedCell."""
    cell = ws.cell(row=r, column=c)
    if cell.__class__.__name__ != 'MergedCell':
        cell.value = val

def copy_row_style(ws, source_row, target_row, min_col=1, max_col=14):
    """
    Copies cell formatting (font, border, fill, number_format, alignment)
    from source_row to target_row across columns min_col..max_col.
    """
    for col in range(min_col, max_col + 1):
        src_cell = ws.cell(row=source_row, column=col)
        dst_cell = ws.cell(row=target_row, column=col)
        if src_cell.has_style:
            dst_cell.font = copy(src_cell.font)
            dst_cell.border = copy(src_cell.border)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.number_format = copy(src_cell.number_format)
            dst_cell.alignment = copy(src_cell.alignment)

def generate_travel_expense_xlsx(expense_id, conn=None):
    """
    Generates a travel expense report in .xlsx format using openpyxl.
    Populates and formats all 5 sheets:
      1. ใบสำคัญจ่ายเงินสดย่อย
      2. รายงาน
      3. คิดเงิน
      4. แบบ บก111
      5. แบบฟอร์ม1
    """
    import sqlite3
    
    close_conn = False
    if conn is None:
        conn = sqlite3.connect(Config.DATABASE_PATH)
        conn.row_factory = sqlite3.Row
        close_conn = True
        
    try:
        # Fetch travel expense details joining staff, positions, sections, divisions, vehicles
        expense_row = conn.execute("""
            SELECT e.id, e.approval_document_number, e.depart_date, e.depart_time,
                   e.return_date, e.return_time, e.fuel_cost, e.toll_cost, e.total_amount,
                   e.created_at,
                   s.full_name as staff_name, s.id as staff_id, s.grade, s.phone, s.position_id,
                   p.position_name,
                   sec.section_name, div.division_name,
                   v.plate_number as vehicle_plate, vt.type_name as vehicle_type,
                   vt.rate_per_km
            FROM travel_expenses e
            LEFT JOIN staff s ON e.staff_id = s.id
            LEFT JOIN positions p ON s.position_id = p.id
            LEFT JOIN sections sec ON s.section_id = sec.id
            LEFT JOIN divisions div ON sec.division_id = div.id
            LEFT JOIN vehicles v ON e.vehicle_id = v.id
            LEFT JOIN vehicle_types vt ON v.vehicle_type_id = vt.id
            WHERE e.id = ?
        """, (expense_id,)).fetchone()
        
        if not expense_row:
            raise ValueError(f"Travel expense record #{expense_id} not found.")
            
        # Fetch expense rates for staff position
        rates_row = conn.execute("""
            SELECT per_diem_rate, accommodation_rate
            FROM expense_rates
            WHERE position_id = ?
            ORDER BY effective_date DESC LIMIT 1
        """, (expense_row['position_id'],)).fetchone()
        
        per_diem_rate = rates_row['per_diem_rate'] if rates_row else 240.0
        accommodation_rate = rates_row['accommodation_rate'] if rates_row else 1200.0
        per_diem_partial_rate = math.floor(per_diem_rate * 2.0 / 3.0)

        # Fetch expense legs
        legs_rows = conn.execute("""
            SELECT leg_order, from_location, depart_date, depart_time,
                   to_location, arrive_date, arrive_time, work_report,
                   distance_km, per_diem_multiplier, has_accommodation
            FROM travel_expense_legs
            WHERE travel_expense_id = ?
            ORDER BY leg_order ASC
        """, (expense_id,)).fetchall()
        
    finally:
        if close_conn:
            conn.close()

    # Define paths
    template_path = os.path.join(Config.TEMPLATE_DOCX_DIR, 'travel_expense_template.xlsx')
    output_filename = f"travel_expense_{expense_id}.xlsx"
    output_path = os.path.join(Config.OUTPUT_DOCX_DIR, output_filename)
    
    # Ensure template file exists
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Excel template not found at {template_path}")
        
    # Ensure output directory exists
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    # Copy template to output destination
    shutil.copy(template_path, output_path)
    
    # Load workbook
    wb = openpyxl.load_workbook(output_path, data_only=False)
    
    # Map sheets by unicode name (handling potential encoding mismatch)
    sheet_map = {}
    for name in wb.sheetnames:
        if "บก111" in name or "111" in name:
            sheet_map['บก111'] = name
        elif "รายงาน" in name or "§ҹ" in name:
            sheet_map['รายงาน'] = name
        elif "ใบสำคัญ" in name or "เงินสดย่อย" in name:
            sheet_map['ใบสำคัญ'] = name
        elif "คิดเงิน" in name:
            sheet_map['คิดเงิน'] = name
        elif "แบบฟอร์ม1" in name or "แบบฟอร์ม" in name:
            sheet_map['แบบฟอร์ม1'] = name
            
    # Fallback to sheet index if mapping failed
    sheet_name_voucher = sheet_map.get('ใบสำคัญ', wb.sheetnames[0] if len(wb.sheetnames) > 0 else 'ใบสำคัญจ่ายเงินสดย่อย')
    sheet_name_rep = sheet_map.get('รายงาน', wb.sheetnames[1] if len(wb.sheetnames) > 1 else 'รายงาน')
    sheet_name_calc = sheet_map.get('คิดเงิน', wb.sheetnames[2] if len(wb.sheetnames) > 2 else 'คิดเงิน')
    sheet_name_bok = sheet_map.get('บก111', wb.sheetnames[3] if len(wb.sheetnames) > 3 else 'แบบ บก111')
    sheet_name_form1 = sheet_map.get('แบบฟอร์ม1', wb.sheetnames[4] if len(wb.sheetnames) > 4 else 'แบบฟอร์ม1')
    
    ws_voucher = wb[sheet_name_voucher]
    ws_rep = wb[sheet_name_rep]
    ws_calc = wb[sheet_name_calc]
    ws_bok = wb[sheet_name_bok]
    ws_form1 = wb[sheet_name_form1]
    
    # Extract data variables
    staff_name = expense_row['staff_name'] or ""
    position_name = expense_row['position_name'] or ""
    grade = expense_row['grade'] or ""
    phone = expense_row['phone'] or ""
    division_name = expense_row['division_name'] or ""
    document_number = expense_row['approval_document_number'] or ""
    fuel_cost = float(expense_row['fuel_cost'] or 0.0)
    toll_cost = float(expense_row['toll_cost'] or 0.0)
    rate_per_km = expense_row['rate_per_km'] or 0.0
    plate_number = expense_row['vehicle_plate'] or ""
    vehicle_type = expense_row['vehicle_type'] or ""
    depart_time = expense_row['depart_time'] or ""
    return_time = expense_row['return_time'] or ""
    
    # Date parsing
    def parse_iso_date(date_str):
        if not date_str:
            return None
        if isinstance(date_str, (datetime.date, datetime.datetime)):
            return date_str
        try:
            return datetime.datetime.strptime(date_str.split(" ")[0], "%Y-%m-%d")
        except ValueError:
            return None
            
    depart_date = parse_iso_date(expense_row['depart_date'])
    return_date = parse_iso_date(expense_row['return_date'])
    document_date = parse_iso_date(expense_row['created_at']) or datetime.date.today()
    if isinstance(document_date, datetime.datetime):
        document_date = document_date.date()

    total_days = 0
    if depart_date and return_date:
        total_days = (return_date - depart_date).days + 1

    # ==========================================
    # 1. SHEET: แบบ บก111 (ใบรับรองแทนใบเสร็จ)
    # ==========================================
    # Clear old cell values in rows 5 to 10
    for r in range(5, 11):
        _set_cell_value(ws_bok, f"A{r}", None)
        _set_cell_value(ws_bok, f"B{r}", None)
        _set_cell_value(ws_bok, f"I{r}", None)
        _set_cell_value(ws_bok, f"J{r}", None)
        
    # Clear draft texts on the side (Column S and rows 29-31)
    _set_cell_value(ws_bok, "S1", None)
    _set_cell_value(ws_bok, "S2", None)
    _set_cell_value(ws_bok, "S3", None)
    for r in range(29, 32):
        _set_cell_value(ws_bok, f"B{r}", None)
        for c in range(19, 26):
            _set_cell_value_rc(ws_bok, r, c, None)
            
    # Set headers in แบบ บก111
    _set_cell_value(ws_bok, "C14", staff_name)
    _set_cell_value(ws_bok, "H14", position_name)
    _set_cell_value(ws_bok, "A15", f"สังกัด {division_name} ขอรับรองว่ารายจ่ายข้างต้นนี้ เป็นค่าใช้จ่ายจริงในงานของ")
    _set_cell_value(ws_bok, "A16", " กปภ. โดยแท้ ")
    
    if depart_date:
        _set_cell_value(ws_bok, "C16", depart_date)
        _set_cell_value(ws_bok, "G16", str(depart_date.year + 543))
    else:
        _set_cell_value(ws_bok, "C16", None)
        _set_cell_value(ws_bok, "G16", None)
        
    if return_date:
        _set_cell_value(ws_bok, "I16", return_date)
        _set_cell_value(ws_bok, "K16", str(return_date.year + 543))
    else:
        _set_cell_value(ws_bok, "I16", None)
        _set_cell_value(ws_bok, "K16", None)
        
    _set_cell_value(ws_bok, "J19", f"({staff_name})")
    _set_cell_value(ws_bok, "J20", position_name)
    
    # Filter legs where vehicle is personal car and distance > 0
    is_personal_car = ("รถยนต์ส่วนตัว" in vehicle_type)
    personal_legs = []
    if is_personal_car:
        personal_legs = [leg for leg in legs_rows if leg['distance_km'] and leg['distance_km'] > 0]
        
    last_row = 4
    for i, leg in enumerate(personal_legs):
        r = 5 + i
        leg_date = parse_iso_date(leg['depart_date'])
        
        _set_cell_value(ws_bok, f"A{r}", leg_date)
        details = (
            f"ค่าเดินทางโดยรถยนต์ส่วนตัวทะเบียน {plate_number} จาก {leg['from_location'] or ''} "
            f"ถึง {leg['to_location'] or ''} ระยะทางไปจำนวน {leg['distance_km'] or 0} "
            f"กิโลเมตร คิด กิโลเมตรละ {rate_per_km} บาท รวมเป็นเงิน"
        )
        _set_cell_value(ws_bok, f"B{r}", details)
        dist = leg['distance_km'] or 0
        _set_cell_value(ws_bok, f"I{r}", f"={dist}*{rate_per_km}")
        _set_cell_value(ws_bok, f"J{r}", "เนื่องจากรถยนต์ กปภ. ไม่เพียงพอ")
        last_row = r
        
    formula_last = max(last_row, 5)
    denominator = rate_per_km if rate_per_km > 0 else 1.0
    _set_cell_value(ws_bok, "B11", f'="รวมระยะทางทั้งหมด "&SUM(I5:I{formula_last})/{denominator}&" กิโลเมตร"')
    _set_cell_value(ws_bok, "I11", f"=SUM(I5:I{formula_last})")
    
    # I12 logic: =I11 if personal car legs exist, otherwise constant 0
    if is_personal_car and len(personal_legs) > 0:
        _set_cell_value(ws_bok, "I12", "=I11")
    else:
        _set_cell_value(ws_bok, "I12", 0)
        
    _set_cell_value(ws_bok, "I13", "=BAHTTEXT(I12)")

    # ==========================================
    # 2. SHEET: คิดเงิน (ใบสรุปค่าใช้จ่ายในการเดินทาง)
    # ==========================================
    def _parse_leg_mult(val):
        if val is None:
            return 0.0
        try:
            return float(val)
        except (ValueError, TypeError):
            return 0.0

    full_day_count = sum(1 for leg in legs_rows if _parse_leg_mult(leg['per_diem_multiplier']) >= 0.99)
    partial_day_count = sum(1 for leg in legs_rows if 0.001 <= _parse_leg_mult(leg['per_diem_multiplier']) < 0.99)
    accommodation_count = sum(1 for leg in legs_rows if leg['has_accommodation'] in [1, True, '1'])

    print(f"[DEBUG XLSX] Travel Expense #{expense_id}: total legs = {len(legs_rows)}")
    for idx, lg in enumerate(legs_rows):
        print(f"  [DEBUG XLSX] Leg {idx+1}: mult={lg['per_diem_multiplier']} (parsed={_parse_leg_mult(lg['per_diem_multiplier'])}), accomm={lg['has_accommodation']}")
    print(f"[DEBUG XLSX] writing C5 (full days) = {full_day_count}")
    print(f"[DEBUG XLSX] writing C6 (partial days) = {partial_day_count}")
    print(f"[DEBUG XLSX] writing C11 (accommodation) = {accommodation_count}")

    _set_cell_value(ws_calc, "A5", staff_name)
    _set_cell_value(ws_calc, "B5", per_diem_rate)
    _set_cell_value(ws_calc, "C5", full_day_count)
    _set_cell_value(ws_calc, "D5", "=B5*C5")

    _set_cell_value(ws_calc, "B6", per_diem_partial_rate)
    _set_cell_value(ws_calc, "C6", partial_day_count)
    _set_cell_value(ws_calc, "D6", "=B6*C6")

    _set_cell_value(ws_calc, "A11", staff_name)
    _set_cell_value(ws_calc, "B11", accommodation_rate)
    _set_cell_value(ws_calc, "C11", accommodation_count)
    _set_cell_value(ws_calc, "D11", "=B11*C11")

    _set_cell_value(ws_calc, "D16", toll_cost)
    _set_cell_value(ws_calc, "D18", fuel_cost)
    _set_cell_value(ws_calc, "C29", f"({staff_name})")
    _set_cell_value(ws_calc, "C30", position_name)

    # ==========================================
    # 3. SHEET: รายงาน (รายงานการเดินทาง)
    # ==========================================
    # Headers in รายงาน
    _set_cell_value(ws_rep, "D4", staff_name)
    _set_cell_value(ws_rep, "H4", position_name)
    
    if document_date:
        ws_rep["H3"].number_format = '0'
        _set_cell_value(ws_rep, "H3", document_date.day)
        _set_cell_value(ws_rep, "J3", thai_months[document_date.month])
        ws_rep["L3"].number_format = '0'
        _set_cell_value(ws_rep, "L3", document_date.year + 543)
    else:
        _set_cell_value(ws_rep, "H3", None)
        _set_cell_value(ws_rep, "J3", None)
        _set_cell_value(ws_rep, "L3", None)

    if depart_date:
        _set_cell_value(ws_rep, "D8", str(depart_date.day))
        _set_cell_value(ws_rep, "F8", thai_months[depart_date.month])
        _set_cell_value(ws_rep, "I8", str(depart_date.year + 543))
    else:
        _set_cell_value(ws_rep, "D8", None)
        _set_cell_value(ws_rep, "F8", None)
        _set_cell_value(ws_rep, "I8", None)
        
    _set_cell_value(ws_rep, "K8", depart_time)
    _set_cell_value(ws_rep, "D6", document_number)
    
    if return_date:
        _set_cell_value(ws_rep, "D9", str(return_date.day))
        _set_cell_value(ws_rep, "F9", thai_months[return_date.month])
        _set_cell_value(ws_rep, "I9", str(return_date.year + 543))
    else:
        _set_cell_value(ws_rep, "D9", None)
        _set_cell_value(ws_rep, "F9", None)
        _set_cell_value(ws_rep, "I9", None)
        
    _set_cell_value(ws_rep, "K9", return_time)
    _set_cell_value(ws_rep, "D10", total_days)

    # Dynamic row handling for table (template default has 4 rows: 13, 14, 15, 16)
    n_legs = len(legs_rows)
    if n_legs > 4:
        extra_rows = n_legs - 4
        ws_rep.insert_rows(17, amount=extra_rows)
        for r in range(17, 17 + extra_rows):
            copy_row_style(ws_rep, 16, r, 1, 14)
            if 16 in ws_rep.row_dimensions and ws_rep.row_dimensions[16].height:
                ws_rep.row_dimensions[r].height = ws_rep.row_dimensions[16].height
            # Merge columns for newly inserted table rows
            ws_rep.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2) # A:B
            ws_rep.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6) # E:F
            ws_rep.merge_cells(start_row=r, start_column=9, end_row=r, end_column=12) # I:L
    elif n_legs < 4:
        # Clear extra rows in template (up to row 16) while keeping borders and styling
        for r in range(13 + n_legs, 17):
            _set_cell_value_rc(ws_rep, r, 1, None)
            _set_cell_value_rc(ws_rep, r, 3, None)
            _set_cell_value_rc(ws_rep, r, 4, None)
            _set_cell_value_rc(ws_rep, r, 5, None)
            _set_cell_value_rc(ws_rep, r, 7, None)
            _set_cell_value_rc(ws_rep, r, 8, None)
            _set_cell_value_rc(ws_rep, r, 9, None)

    # Write legs data to รายงาน
    for i, leg in enumerate(legs_rows):
        r = 13 + i
        leg_dept_date = parse_iso_date(leg['depart_date'])
        leg_arr_date = parse_iso_date(leg['arrive_date'])
        
        dept_date_str = to_thai_date_excel_short(leg_dept_date)
        arr_date_str = to_thai_date_excel_short(leg_arr_date)
        
        base_text = leg['work_report'] or ""
        if vehicle_type and ("รถเช่า" in vehicle_type):
            work_report_val = base_text + f" โดยรถเช่า กปภ. ทะเบียน {plate_number}"
        elif vehicle_type == "รถยนต์ส่วนตัว":
            work_report_val = base_text + f" โดยรถยนต์ส่วนตัว ทะเบียน {plate_number}"
        else:
            work_report_val = base_text
            
        _set_cell_value(ws_rep, f"A{r}", leg['from_location'] or "")
        _set_cell_value(ws_rep, f"C{r}", dept_date_str)
        _set_cell_value(ws_rep, f"D{r}", leg['depart_time'] or "")
        _set_cell_value(ws_rep, f"E{r}", leg['to_location'] or "")
        _set_cell_value(ws_rep, f"G{r}", arr_date_str)
        _set_cell_value(ws_rep, f"H{r}", leg['arrive_time'] or "")
        _set_cell_value(ws_rep, f"I{r}", work_report_val)

    # ==========================================
    # 4. SHEET: ใบสำคัญจ่ายเงินสดย่อย
    # ==========================================
    _set_cell_value(ws_voucher, "G6", "X")
    _set_cell_value(ws_voucher, "T11", "=SUM(คิดเงิน!D5:D6)")
    _set_cell_value(ws_voucher, "T12", "=คิดเงิน!D11")
    _set_cell_value(ws_voucher, "T13", fuel_cost)
    _set_cell_value(ws_voucher, "T14", "=คิดเงิน!D16")
    _set_cell_value(ws_voucher, "T15", "=คิดเงิน!D17")
    _set_cell_value(ws_voucher, "T17", "=SUM(T11:T15)")

    pos_and_grade = f"{position_name} {grade}".strip()
    _set_cell_value(ws_voucher, "C24", pos_and_grade)
    _set_cell_value(ws_voucher, "C41", pos_and_grade)
    _set_cell_value(ws_voucher, "H24", document_date)
    _set_cell_value(ws_voucher, "H41", document_date)

    # ==========================================
    # 5. SHEET: แบบฟอร์ม1 (ใบแจ้งค่าใช้จ่ายค้างจ่าย)
    # ==========================================
    _set_cell_value(ws_form1, "B5", staff_name)
    _set_cell_value(ws_form1, "H5", to_thai_date_full(document_date))
    _set_cell_value(ws_form1, "H6", phone)
    _set_cell_value(ws_form1, "K6", thai_months[document_date.month])
    _set_cell_value(ws_form1, "L6", f"ปี {document_date.year + 543}")
    _set_cell_value(ws_form1, "B9", f"ค่าใช้จ่ายในการเดินทาง {staff_name}")
    _set_cell_value(ws_form1, "F9", staff_name)
    _set_cell_value(ws_form1, "J9", "='ใบสำคัญจ่ายเงินสดย่อย'!T17")
    _set_cell_value(ws_form1, "L9", "='ใบสำคัญจ่ายเงินสดย่อย'!T17")
    
    # Save workbook
    wb.save(output_path)
    return output_filename
